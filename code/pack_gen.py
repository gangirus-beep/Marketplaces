# Упаковка /home/claude/gen в «Генератор_обложек_WB.xlsx»: python3 pack_gen.py "<строка версии>"
import os, io, zipfile, base64, hashlib, openpyxl, sys
from openpyxl.styles import Font, Alignment
root='/home/claude/gen'; files=[]
for dp,dn,fn in os.walk(root):
    rel=os.path.relpath(dp,root)
    if rel!='.' and rel.split(os.sep)[0] not in {'fonts','pass','pilot','stage'}: continue
    for f in fn:
        if rel=='.' and not (f=='axis.py' or f.startswith('catalog_') or f in ('README.md','cover.html','pack_gen.py')): continue
        files.append(os.path.normpath(os.path.join(rel,f)))
files.sort(); buf=io.BytesIO()
with zipfile.ZipFile(buf,'w',zipfile.ZIP_DEFLATED,compresslevel=9) as z:
    for f in files: z.write(os.path.join(root,f),f)
raw=buf.getvalue(); sha=hashlib.sha256(raw).hexdigest(); b64=base64.b64encode(raw).decode()
code="import base64, io, zipfile, openpyxl\nwb = openpyxl.load_workbook('/mnt/project/Генератор_обложек_WB.xlsx', read_only=True)\ndata = ''.join(r[1] for r in wb['Файлы'].iter_rows(min_row=2, values_only=True) if r[1])\nzipfile.ZipFile(io.BytesIO(base64.b64decode(data))).extractall('/home/claude/gen')\n# дальше — README.md в распакованной папке; рабочая папка этапа — /home/claude/gen"
wb=openpyxl.Workbook(); ws=wb.active; ws.title='Как пользоваться'; ws.column_dimensions['A'].width=150
rows=[sys.argv[1],
 'Этот файл не нужно открывать и править: он лежит в файлах проекта для Claude. Проект не принимает zip-архивы, поэтому генератор упакован в лист «Файлы».',
 'Внутри: шаблон обложки cover.html со всеми фонами, шрифты с открытой лицензией, папка pass — проход «в упаковке / поштучно» (скачивание фото типа, миниатюры и поиск одинаковых изображений, контактные листы, крупный просмотр, обрезка слайдов Рельефа по контуру товара, два образца решений (этапы 2 и 3), запись колонок U–AC в базу и сверка), папка stage — конвейер одного этапа (выборка из базы, подготовка фото, загрузка кадра с обрезкой слайдов по контуру, вырезка, вырезка «по фигуре» из прямых линий с заходом в товар (в том числе ступенчатая фигура и пять прямоугольников для коллажей внахлёст), повторная вырезка построчно и по оболочке, три образца текстов (этапы 1, 2 и 3), веер, рендер обложек, последние фото, листы просмотра, дозапись реестра, архив, перевыпуск обложек: обновление реестра и архив только с обложками), папка pilot — скрипты пилота, README с порядком работы, pack_gen.py — упаковка генератора обратно в этот файл.',
 'Состав: '+', '.join(files),'Контрольная сумма архива (sha256): '+sha,'Код распаковки для Claude:',code]
for i,t in enumerate(rows,1):
    c=ws.cell(i,1,t); c.font=Font(name='Courier New' if i==7 else 'Arial',bold=(i==1)); c.alignment=Alignment(wrap_text=True,vertical='top')
w2=wb.create_sheet('Файлы'); w2.column_dimensions['B'].width=60; w2.append(['Часть','Данные (base64 zip-архива)'])
for c in w2[1]: c.font=Font(name='Arial',bold=True)
for i in range(0,len(b64),30000): w2.append([str(i//30000+1),b64[i:i+30000]])
out='/mnt/user-data/outputs/Генератор_обложек_WB.xlsx'; wb.save(out)
data=''.join(r[1] for r in openpyxl.load_workbook(out,read_only=True)['Файлы'].iter_rows(min_row=2,values_only=True) if r[1])
r2=base64.b64decode(data); assert hashlib.sha256(r2).hexdigest()==sha and zipfile.ZipFile(io.BytesIO(r2)).testzip() is None
print(len(files),'files, xlsx KB',os.path.getsize(out)//1024,sha[:12])
