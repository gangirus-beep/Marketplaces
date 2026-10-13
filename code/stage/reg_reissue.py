# Перевыпуск обложек этапа (решение владельца 06.10.2026: обложки пилота и этапов 1–2 перевыпускаются с вырезкой «по фигуре»).
# Строки этапа в реестре не дописываются, а обновляются: новая дата, шрифт и тон из coverlog.json, в колонке «Архив» — архив перевыпуска
# для обложки и прежний архив для последнего фото; на листе «Этапы» — статус и дополнение к комментарию; новые строки листа «Правила».
# Параметры — в texts.py: STAGE, STAGE_NO, TYPE, ARCH (прежний архив) и блок перевыпуска REDATE, REARCH, RECOMMENT, RE_RULES.
# Перед записью сверяется, что набор файлов обложек совпадает с реестром один к одному и тексты не изменились.
import json, openpyxl
from openpyxl.styles import Font, Alignment
import texts
SRC='/mnt/project/Реестр_обложек_и_последних_фото.xlsx'; DST='/mnt/user-data/outputs/Реестр_обложек_и_последних_фото.xlsx'
goods=json.load(open('goods.json')); meta=json.load(open('cutmeta.json')); clog={l['file']:l for l in json.load(open('coverlog.json'))}
BF=Font(name='Arial',size=10); AL=Alignment(vertical='top',wrap_text=True)
wb=openpyxl.load_workbook(SRC); ws=wb['Готово']; H=[c.value for c in ws[1]]; I={n:i for i,n in enumerate(H)}
rows=[r for r in ws.iter_rows(min_row=2) if r[I['Этап']].value==texts.STAGE]
files={r[I['Файл обложки']].value for r in rows}; assert files==set(clog), (len(files),len(clog),sorted(files^set(clog))[:5])
code={x['code']:x for x in goods}; ch=dict(font=0,tone=0)
for r in rows:
    f=r[I['Файл обложки']].value; l=clog[f]; x=code[str(r[I['Код поставщика']].value)]
    assert r[I['Архив']].value==texts.ARCH, r[I['Архив']].value
    font={'grotesk':'гротеск','serif':'с засечками'}.get(l['font'],'узкий (длинный заголовок)'); tone='тёмная подложка под фото' if l['tone']=='panel' else 'светлый'
    ch['font']+=r[I['Шрифт заголовка']].value!=font; ch['tone']+=r[I['Тон обложки']].value!=tone
    r[I['Дата']].value=texts.REDATE; r[I['Шрифт заголовка']].value=font; r[I['Тон обложки']].value=tone
    r[I['Архив']].value=('обложка — '+texts.REARCH+'; последнее фото — '+texts.ARCH) if r[I['Файл последнего фото']].value!='НЕТ' else texts.REARCH
print('Готово: обновлено строк',len(rows),'| изменился шрифт',ch['font'],'| тон',ch['tone'])
ws=wb['Тексты обложек']; n=0
for r in ws.iter_rows(min_row=2,values_only=True):
    if r[0]==texts.STAGE:
        t=texts.T[str(r[2])]; assert (r[6],r[7],list(r[8:12]))==(t['heading'],t['sub'],t['chips']), r[2]; n+=1
print('Тексты обложек: сверено',n,'— без изменений')
ws=wb['Этапы']
for r in ws.iter_rows(min_row=2):
    if r[0].value==texts.STAGE_NO and r[3].value==texts.TYPE and (r[4].value or None)==getattr(texts,'STAGE_PART',None):
        r[1].value=r[1].value.split(' · ')[0]+' · обложки перевыпущены '+texts.REDATE; r[12].value=(r[12].value or '').split(' Перевыпуск обложек ')[0]+' '+texts.RECOMMENT; r[12].font=BF; r[12].alignment=AL; print('Этапы:',r[1].value)
ws=wb['Правила']
for t in getattr(texts,'RE_RULES',[]):
    ws.append([t]); c=ws.cell(ws.max_row,1); c.font=BF; c.alignment=AL
wb.save(DST); print('saved')
