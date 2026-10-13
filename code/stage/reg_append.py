# Шаг 7. Дозапись этапа в накопительный реестр: листы «Готово», «Тексты обложек», статус на листе «Этапы», новые правила
import json, copy, collections, openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
import texts
from last import layname
SRC='/mnt/project/Реестр_обложек_и_последних_фото.xlsx'; DST='/mnt/user-data/outputs/Реестр_обложек_и_последних_фото.xlsx'
DATE=texts.DATE; STAGE=texts.STAGE; TYPE=texts.TYPE; ARCH=texts.ARCH     # параметры этапа — в texts.py
goods=json.load(open('goods.json')); meta=json.load(open('cutmeta.json')); last={l['file']:l for l in json.load(open('lastlog.json'))}; clog={l['file']:l for l in json.load(open('coverlog.json'))}
BF=Font(name='Arial',size=10); AL=Alignment(vertical='top',wrap_text=True)
def forms(q,f):
    f=f.split('|'); n10,n100=q%10,q%100
    return f[0] if (n10==1 and n100!=11) else f[1] if (2<=n10<=4 and not 12<=n100<=14) else f[2]
wb=openpyxl.load_workbook(SRC)
def add(ws,rows):
    for r in rows:
        ws.append(r)
        for c in ws[ws.max_row]: c.font=BF; c.alignment=AL
    ws.auto_filter.ref=ws.dimensions
ws=wb['Готово']; done={r[11] for r in ws.iter_rows(min_row=2,values_only=True)}   # файлы обложек, уже записанные в реестр
rows=[]
for x in goods:
    t=texts.T[x['code']]; tone=texts.TONE_FORCE.get(x['code'],meta[x['code']]['tone']); single=all(k['quant']==1 for k in x['cards'])
    for c in x['cards']:
        q=c['quant']; cf=f"{c['idvh']}_{q}.jpg"; lf=f"{c['idvh']}_{q}_last.jpg"; L=last.get(lf); cl=clog[cf]
        nt=texts.NOTE.get(x['code'],'')
        if L and q>=100 and getattr(texts,'SMALL_NOTE',True): nt=(nt+f' Квант {q}: на последнем фото единицы очень мелкие ({L["item_w"]} px по ширине) — посмотреть.').strip()
        rows.append([DATE,STAGE,TYPE,int(c['wb']),c['seller'],int(c['idvh']),q,x['sup'],x['code'],x['name'],x['status'],cf,lf if L else 'НЕТ',texts.BGCODE[t['bg']],texts.PLAQUE,
            'гротеск' if cl['font']=='grotesk' else 'узкий (длинный заголовок)','тёмная подложка под фото' if tone=='panel' else 'светлый',
            'без плашки' if single else f"{q} {forms(q,t['forms'])}"+((' / '+t['line2_1'] if t.get('line2_1') else '') if q==1 else ' '+t.get('line2',('в комплекте' if x['unit'] in ('набор','упаковка') else 'в наборе'))),f"№{x['U']}" + (' · обрезка разметки по контуру' if 'slidecrop.py' in (x['comment'] or '') else '') + ((' · коллаж веером, один раз' if t.get('fan') else ' · коллаж, один раз') if t.get('copies') else ''),(f"№{x['W']}" + (' · обрезка разметки по контуру' if 'slidecrop.py' in (x['comment'] or '') else '')) if L else '—',
            L['layout'] if L else 'нет: «ассорти», на фото варианты дизайна',nt,ARCH])
dup=[r[11] for r in rows if r[11] in done]; assert not dup, ('уже в реестре',dup[:5])
add(ws,rows); print('Готово +',len(rows),'→',ws.max_row-1)
ws=wb['Тексты обложек']; rows=[]
for x in goods:
    t=texts.T[x['code']]; ids=list(dict.fromkeys(c['idvh'] for c in x['cards']))
    rows.append([STAGE,int(ids[0]) if len(ids)==1 else ', '.join(ids),x['code'],x['name'],', '.join(str(c['quant']) for c in x['cards']),t['forms'],t['heading'],t['sub']]+t['chips']+[texts.NOTE.get(x['code'],'')])
add(ws,rows); print('Тексты +',len(rows))
ws=wb['Этапы']
for r in ws.iter_rows(min_row=2):
    if r[0].value==texts.STAGE_NO and r[3].value==TYPE and (r[4].value or None)==getattr(texts,'STAGE_PART',None):
        r[1].value='сделано '+DATE; r[12].value=texts.STAGE_COMMENT; r[12].font=BF; r[12].alignment=AL; print('Этапы: статус записан')
ws=wb['Правила']
for r in getattr(texts,'NEW_RULES',[]):
    ws.append([r]); c=ws.cell(ws.max_row,1); c.font=BF; c.alignment=AL
wb.save(DST); print('saved')
